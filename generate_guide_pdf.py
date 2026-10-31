import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Set inner padding for table cell (in twips: 20 twips = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_callout_borders(cell, border_color_hex, border_size="24"):
    """Set thick left border and none for others."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="{border_size}" w:space="0" w:color="{border_color_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_table_borders(table, border_color="CBD5E0"):
    """Set subtle light gray borders for a standard data table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{border_color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_callout(doc, title, text, callout_type="action"):
    """
    Adds a styled callout box with custom background and left border.
    Types: 'action', 'screenshot', 'tip', 'warning'
    """
    styles = {
        "action": {
            "bg": "FFFBEB",
            "border": "D69E2E",
            "icon": "👉 [HÀNH ĐỘNG DÀNH CHO BẠN]",
            "title_color": RGBColor(183, 121, 31)
        },
        "screenshot": {
            "bg": "EBF8FF",
            "border": "3182CE",
            "icon": "📸 [NHẮC NHỞ CHỤP MÀN HÌNH ĐỂ LƯU MINH CHỨNG]",
            "title_color": RGBColor(43, 108, 176)
        },
        "tip": {
            "bg": "F0FFF4",
            "border": "38A169",
            "icon": "💡 [MẸO CHUYÊN GIA DÀNH CHO NGƯỜI MỚI]",
            "title_color": RGBColor(47, 133, 90)
        },
        "warning": {
            "bg": "FFF5F5",
            "border": "E53E3E",
            "icon": "⚠️ [LƯU Ý QUAN TRỌNG - KHÔNG ĐƯỢC BỎ QUA]",
            "title_color": RGBColor(229, 62, 62)
        }
    }
    
    cfg = styles.get(callout_type, styles["action"])
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_shading(cell, cfg["bg"])
    set_callout_borders(cell, cfg["border"], border_size="28")
    
    # Title paragraph
    p_title = cell.paragraphs[0]
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.line_spacing = 1.15
    run_icon = p_title.add_run(f"{cfg['icon']}\n")
    run_icon.bold = True
    run_icon.font.name = "Calibri"
    run_icon.font.size = Pt(10.5)
    run_icon.font.color.rgb = cfg["title_color"]
    
    run_sub = p_title.add_run(title)
    run_sub.bold = True
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(45, 55, 72)
    
    # Text paragraph
    p_text = cell.add_paragraph()
    p_text.paragraph_format.space_before = Pt(2)
    p_text.paragraph_format.space_after = Pt(2)
    p_text.paragraph_format.line_spacing = 1.2
    run_text = p_text.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(45, 55, 72)
    
    # Empty space after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, caption=None):
    """Adds a stylish terminal/code snippet block."""
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"📋 {caption}")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(74, 85, 104)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_shading(cell, "F7FAFC")
    set_callout_borders(cell, "CBD5E0", border_size="16")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(45, 55, 72)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def build_pdf_document():
    doc = Document()
    
    # Configure page geometry (A4, 0.8 inch margins for great layout)
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)  # A4 width
        section.page_height = Inches(11.69) # A4 height
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
        # Header and Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("Hướng Dẫn MLOps: Triển Khai CI/CD & Public Docker Hub Cho Mô Hình AI")
        hr.font.name = "Calibri"
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(160, 174, 192)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("Dự án AI: Music_CoTMix • Tài liệu Hướng Dẫn Từng Bước Dành Cho Người Không Chuyên")
        fr.font.name = "Calibri"
        fr.font.size = Pt(8.5)
        fr.font.color.rgb = RGBColor(160, 174, 192)

    # Helper styling functions
    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(22)
        r.bold = True
        r.font.color.rgb = RGBColor(26, 54, 93) # Deep Navy
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11.5)
        r.font.italic = True
        r.font.color.rgb = RGBColor(74, 85, 104)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(15)
        r.bold = True
        r.font.color.rgb = RGBColor(43, 108, 176) # Accent Blue
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(12.5)
        r.bold = True
        r.font.color.rgb = RGBColor(44, 82, 130) # Medium Blue
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.bold = True
        r.font.color.rgb = RGBColor(45, 55, 72)
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10.5)
            rb.bold = True
            rb.font.color.rgb = RGBColor(45, 55, 72)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(45, 55, 72)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Calibri"
            rb.font.size = Pt(10)
            rb.bold = True
            rb.font.color.rgb = RGBColor(45, 55, 72)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(45, 55, 72)
        return p

    # -------------------------------------------------------------
    # DOCUMENT CONTENT
    # -------------------------------------------------------------
    
    # Title & Metadata Card
    add_title("CẨM NANG HƯỚNG DẪN TRIỂN KHAI CI/CD & PUBLIC DOCKER HUB CHO BÀI TOÁN AI")
    add_subtitle("Tài liệu hướng dẫn trực quan, từng bước một (Step-by-Step) dành cho người không chuyên về kỹ thuật\nDự án: Hệ Thống Nhận Diện Thể Loại Âm Nhạc Trực Tuyến (Music_CoTMix)")

    # Metadata Summary Box
    meta_tbl = doc.add_table(rows=4, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Mục tiêu bài toán:", "1. Tự động hóa kiểm thử & triển khai (CI/CD) qua GitHub Actions\n2. Đóng gói mô hình AI vào Docker Image và xuất bản (Public) lên Docker Hub"),
        ("Đối tượng tiếp nhận:", "Người quản lý, sinh viên, người mới bắt đầu (không yêu cầu kinh nghiệm DevOps/Docker)"),
        ("Mã nguồn dự án:", "Music_CoTMix (FastAPI + PyTorch AudioCNN/DinoV3 Model)"),
        ("Thời gian thực hiện:", "Chỉ mất từ 10 - 15 phút với các bước click chuột được chỉ dẫn chi tiết")
    ]
    set_table_borders(meta_tbl, "CBD5E0")
    for idx, (label, val) in enumerate(meta_info):
        c0 = meta_tbl.cell(idx, 0)
        c1 = meta_tbl.cell(idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_margins(c0, 80, 80, 120, 120)
        set_cell_margins(c1, 80, 80, 120, 120)
        set_cell_shading(c0, "EDF2F7")
        set_cell_shading(c1, "F7FAFC")
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.bold = True
        r0.font.color.rgb = RGBColor(43, 108, 176)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(45, 55, 72)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # SECTION 1
    add_h1("PHẦN 1: GIẢI THÍCH KHÁI NIỆM DỄ HIỂU NHẤT (VÍ DỤ ĐỜI THƯỜNG)")
    add_p(
        "Đối với một người chưa từng học lập trình hay quản trị hệ thống, các từ ngữ như 'CI/CD', 'Docker', "
        "'Container', 'Docker Hub' nghe có vẻ rất phức tạp và bí ẩn. Tuy nhiên, bản chất của chúng lại vô cùng gần gũi với đời sống thực tế:"
    )

    add_bullet(
        " Giống như một chiếc thùng container tiêu chuẩn bằng kim loại dùng trong vận chuyển đường biển quốc tế. "
        "Trước đây, mỗi lần chuyển một mô hình AI sang máy tính khác, chúng ta rất dễ gặp lỗi do thiếu thư viện âm thanh, sai phiên bản Python, "
        "hoặc hệ điều hành không tương thích ('ở máy tôi chạy tốt nhưng máy bạn thì lỗi'). Docker giải quyết triệt để điều này bằng cách đóng gói "
        "toàn bộ: hệ điều hành thu nhỏ, mã nguồn FastAPI, file trọng số mô hình AI (.pt) và toàn bộ thư viện cần thiết vào một khối thống nhất. "
        "Bất kỳ máy tính nào có Docker cũng sẽ mở thùng này và chạy đúng 100% như mong đợi.",
        "1. Docker Image là gì? (Chiếc thùng Container chở hàng):"
    )

    add_bullet(
        " Giống hệt như cửa hàng ứng dụng Google Play Store hoặc Apple App Store trên điện thoại của bạn. "
        "Thay vì chia sẻ qua Zalo hay Google Drive rất nặng và dễ lỗi, Docker Hub là một website toàn cầu cho phép bạn đưa 'chiếc thùng container' "
        "của mình lên đó (Public). Bất kỳ ai trên thế giới cũng có thể chỉ cần gõ 1 dòng lệnh là tự động tải mô hình của bạn về dùng ngay lập tức.",
        "2. Docker Hub là gì? (Kho ứng dụng App Store dành cho Container):"
    )

    add_bullet(
        " Hãy tưởng tượng bạn mở một xưởng sản xuất ô tô. Khi một kỹ sư làm xong bộ phanh xe mới:\n"
        "• CI (Continuous Integration - Tích hợp liên tục): Dây chuyền tự động kích hoạt để kiểm tra xem phanh có ăn không, có bị gãy không (chạy các bài test tự động).\n"
        "• CD (Continuous Deployment - Phân phối liên tục): Sau khi phanh xe đạt tiêu chuẩn kiểm nghiệm, robot tự động sơn bóng, đóng thùng container và chuyển thẳng lên tàu biển (đưa lên Docker Hub hoặc máy chủ đám mây).\n"
        "Toàn bộ việc này diễn ra 100% tự động, bạn không cần phải ngồi bấm từng lệnh phức tạp trên máy tính của mình!",
        "3. CI/CD là gì? (Dây chuyền kiểm định và giao hàng tự động):"
    )

    add_bullet(
        " Đóng vai trò là 'anh công nhân robot miễn phí' của GitHub. Mỗi khi bạn tải mã nguồn mới lên GitHub, "
        "GitHub Actions sẽ lập tức khởi động một máy chủ ảo trên mây của Microsoft, tự chạy kiểm tra mã nguồn, tự đóng gói Docker và tự đưa lên Docker Hub cho bạn.",
        "4. GitHub Actions là gì? (Người máy phụ tá trên mây):"
    )

    add_callout(
        doc,
        "Tại sao bạn không cần phải cài đặt phần mềm nặng nhọc trên máy cá nhân?",
        "Vì máy tính của bạn hiện tại chưa cài Docker Desktop (vốn khá nặng và chiếm nhiều RAM), chúng tôi đã thiết lập sẵn toàn bộ quy trình CI/CD trên GitHub Actions. "
        "Điều đó có nghĩa là: Bạn KHÔNG CẦN phải cài Docker trên máy tính của mình! Toàn bộ việc build, đóng gói và public lên Docker Hub sẽ do máy chủ của GitHub thực hiện thay bạn hoàn toàn miễn phí!",
        "tip"
    )

    # SECTION 2
    add_h1("PHẦN 2: NHỮNG FILE KỸ THUẬT ĐÃ ĐƯỢC CHUYÊN GIA THIẾT LẬP SẴN")
    add_p("Trong thư mục dự án của bạn, chuyên gia phát triển phần mềm đã tạo và cấu hình hoàn chỉnh các thành phần sau:")

    add_bullet(" Tệp chỉ dẫn các bước đóng gói mô hình AI: tải Python 3.11 siêu nhẹ, cài các thư viện âm thanh (libsndfile, ffmpeg), nạp trọng số mô hình Music_CoTMix và kích hoạt API dự đoán thể loại nhạc.", "• Dockerfile:")
    add_bullet(" Bộ lọc thông minh ngăn chặn những file thừa (như 1.8GB nhạc wav thừa, các tệp ghi chú jupyter notebook, file tạm) không bị copy vào Docker, giúp dung lượng container giảm từ hơn 3.5GB xuống mức siêu nhẹ chỉ còn vài trăm MB.", "• .dockerignore:")
    add_bullet(" Kịch bản kiểm tra tự động trước khi xuất xưởng: kiểm tra cổng mạng /health, kiểm tra việc nạp mô hình vào bộ nhớ, kiểm tra luồng dự đoán thể loại bài hát.", "• tests/test_api.py:")
    add_bullet(" Quy trình CI tự động kích hoạt trên GitHub Actions: kiểm tra cú pháp mã nguồn (Linting) và chạy kiểm thử tự động (Unit Test).", "• .github/workflows/ci.yml:")
    add_bullet(" Quy trình CD tự động kích hoạt trên GitHub Actions: tự động đăng nhập Docker Hub, tự động đóng gói Docker Image đa nền tảng và tải công khai lên Docker Hub.", "• .github/workflows/cd-docker.yml:")
    add_bullet(" Công cụ 1-click dành cho bạn nếu muốn tự build trên máy tính Windows cá nhân khi đã cài Docker Desktop.", "• docker_build_push.bat & docker-compose.yml:")

    # SECTION 3
    add_h1("PHẦN 3: HƯỚNG DẪN TỪNG BƯỚC CHO BẠN (CỰC KỲ DỄ THỰC HIỆN)")
    add_p("Dưới đây là phần việc duy nhất mà bạn cần thực hiện. Mọi thao tác đều là click chuột trên trình duyệt web, không đòi hỏi viết code:")

    add_h2("Bước 1: Đăng Ký Tài Khoản Docker Hub (Miễn Phí)")
    add_p("Docker Hub là nơi chứa Docker Image công khai của bạn để mọi người có thể xem và tải về.")
    add_bullet("Mở trình duyệt web (Chrome, Edge, Cốc Cốc) và truy cập địa chỉ: https://hub.docker.com", "1.1. ")
    add_bullet("Bấm vào nút 'Sign Up' màu xanh ở góc trên bên phải màn hình.", "1.2. ")
    add_bullet("Nhập thông tin đăng ký: Docker ID (Tên đăng nhập của bạn, viết liền không dấu, ví dụ: 'nimesa0' hoặc 'tungdev'), Email và Mật khẩu.", "1.3. ")
    add_bullet("Mở hộp thư Email của bạn để bấm liên kết xác nhận tài khoản (Verify Email).", "1.4. ")

    add_callout(
        doc,
        "Lưu lại tên đăng nhập Docker Hub của bạn",
        "Hãy ghi nhớ chính xác Docker ID (Tên đăng nhập) của bạn. Ví dụ: nimesa0.\nĐây chính là tên thư mục chứa mô hình AI của bạn trên toàn cầu!",
        "action"
    )

    add_callout(
        doc,
        "Chụp lại màn hình trang chủ Docker Hub sau khi đăng nhập thành công",
        "Sau khi đăng nhập vào hub.docker.com, bạn hãy chụp ảnh màn hình trang chủ có hiện tên người dùng của bạn ở góc phải trên cùng để lưu làm minh chứng trong báo cáo nghiệm thu.",
        "screenshot"
    )

    add_h2("Bước 2: Tạo Mật Mã Bảo Mật (Personal Access Token) Trên Docker Hub")
    add_p(
        "Để cho phép robot GitHub Actions thay mặt bạn tải Docker Image lên Docker Hub mà không cần để lộ mật khẩu chính của tài khoản, "
        "bạn cần tạo một chiếc 'chìa khóa phụ' gọi là Access Token:"
    )
    add_bullet("Trên trang web Docker Hub, bấm vào avatar hoặc tên tài khoản của bạn ở góc trên bên phải, chọn 'Account Settings'.", "2.1. ")
    add_bullet("Ở menu bên trái, bấm vào mục 'Security' (Bảo mật).", "2.2. ")
    add_bullet("Bấm vào nút màu xanh 'New Access Token' (Tạo mã truy cập mới).", "2.3. ")
    add_bullet("Điền thông tin vào bảng hiện ra:\n• Description (Mô tả): Gõ chữ 'github-actions-ci'\n• Access permissions (Quyền hạn): Chọn 'Read, Write, Delete' (Đọc, Ghi, Xóa).", "2.4. ")
    add_bullet("Bấm nút 'Generate' (Tạo). Một chuỗi ký tự dài sẽ xuất hiện (bắt đầu bằng chữ dckr_pat_...).", "2.5. ")
    add_bullet("Bấm vào nút 'Copy and Close' để sao chép chuỗi ký tự này và dán tạm vào Notepad.", "2.6. ")

    add_callout(
        doc,
        "Chụp ảnh hộp thoại mã Access Token vừa tạo",
        "Hãy chụp lại ảnh màn hình lúc Docker Hub hiển thị hộp thoại chứa Access Token (lưu ý: bạn có thể che bớt vài ký tự đuôi nếu cần bảo mật, nhưng chụp để làm minh chứng đã tạo token thành công).",
        "screenshot"
    )

    add_callout(
        doc,
        "Access Token chỉ hiển thị duy nhất 1 lần!",
        "Hãy chắc chắn bạn đã bấm nút 'Copy' lưu lại chuỗi token này vào Notepad trước khi đóng cửa sổ. Nếu lỡ quên, bạn sẽ phải bấm tạo lại một token mới.",
        "warning"
    )

    add_h2("Bước 3: Nhập Thông Tin Bí Mật (Secrets) Vào Kho GitHub Của Bạn")
    add_p(
        "Bây giờ chúng ta sẽ giao 'chìa khóa' này cho robot GitHub Actions bằng cách lưu vào két sắt bảo mật (Secrets) của GitHub repository:"
    )
    add_bullet("Mở trình duyệt và truy cập vào dự án GitHub của bạn tại địa chỉ: https://github.com/1Nimesa0/Music_CoTMix", "3.1. ")
    add_bullet("Bấm vào tab 'Settings' (biểu tượng bánh răng) ở thanh điều hướng phía trên của repository.", "3.2. ")
    add_bullet("Ở thanh menu bên trái, tìm mục 'Secrets and variables', bấm vào đó và chọn 'Actions'.", "3.3. ")
    add_bullet("Bấm vào nút màu xanh 'New repository secret'.", "3.4. ")
    add_bullet("Tạo bí mật thứ nhất:\n• Name: Gõ chính xác DOCKER_USERNAME\n• Secret: Nhập tên đăng nhập Docker Hub của bạn (ví dụ: nimesa0)\n• Bấm nút 'Add secret'.", "3.5. ")
    add_bullet("Bấm tiếp 'New repository secret' để tạo bí mật thứ hai:\n• Name: Gõ chính xác DOCKER_PASSWORD\n• Secret: Dán chuỗi ký tự Access Token bạn đã copy ở Bước 2 vào đây\n• Bấm nút 'Add secret'.", "3.6. ")

    add_callout(
        doc,
        "Chụp lại ảnh danh sách Secrets trên GitHub",
        "Sau khi thêm xong, bạn sẽ thấy 2 mục DOCKER_USERNAME và DOCKER_PASSWORD xuất hiện trong bảng 'Repository secrets'. Hãy chụp ảnh màn hình này để chứng minh cấu hình CI/CD đã hoàn tất!",
        "screenshot"
    )

    # SECTION 4
    add_h1("PHẦN 4: KÍCH HOẠT VÀ THEO DÕI QUY TRÌNH CI/CD TỰ ĐỘNG CHẠY")
    add_p("Sau khi đã cấu hình xong 2 Secret ở Bước 3, hệ thống CI/CD đã sẵn sàng 100%. Bạn có 2 cách cực kỳ đơn giản để kích hoạt dây chuyền tự động:")

    add_h2("Cách 1: Đẩy (Push) Các Tệp Mới Lên GitHub (Khuyên Dùng)")
    add_p("Toàn bộ các tệp cấu hình CI/CD và Dockerfile tối ưu đã được chuyên gia tạo sẵn trong thư mục máy tính của bạn. Bạn chỉ cần mở cửa sổ dòng lệnh (Terminal/PowerShell) tại thư mục dự án và chạy 3 câu lệnh sau:")
    add_code_block(
        doc,
        "git add .\n"
        "git commit -m \"feat: Trien khai CI/CD va Docker Hub Publishing\"\n"
        "git push origin main",
        "3 câu lệnh để đưa toàn bộ thiết lập lên GitHub"
    )
    add_p(
        "Ngay khi lệnh 'git push' hoàn tất, GitHub sẽ phát hiện có cập nhật mới và tự động kích hoạt cả 2 dây chuyền CI và CD!"
    )

    add_h2("Cách 2: Bấm Nút Chạy Thủ Công (Run Workflow) Trực Tiếp Trên Web GitHub")
    add_p("Nếu không muốn dùng dòng lệnh, bạn có thể tự kích hoạt bằng cách click chuột:")
    add_bullet("Vào trang repository trên GitHub (https://github.com/1Nimesa0/Music_CoTMix).", "1. ")
    add_bullet("Bấm vào tab 'Actions' ở thanh menu trên cùng.", "2. ")
    add_bullet("Ở cột bên trái, bấm chọn workflow 'AI Model CD Pipeline (Build & Publish to Docker Hub)'.", "3. ")
    add_bullet("Ở bên phải sẽ xuất hiện nút màu xám 'Run workflow'. Bạn bấm vào đó và chọn nhánh 'main', sau đó bấm nút xanh 'Run workflow'.", "4. ")

    add_callout(
        doc,
        "Xem tiến trình chạy tự động (Live Logs) trên GitHub Actions",
        "Bấm vào lượt chạy vừa xuất hiện (có biểu tượng vòng tròn màu vàng đang xoay). Bạn sẽ thấy robot của GitHub đang tự thực hiện từng việc:\n"
        "1. Tải mã nguồn về máy ảo.\n"
        "2. Đăng nhập vào Docker Hub bằng bí mật đã lưu.\n"
        "3. Đóng gói mô hình AI thành Docker Image.\n"
        "4. Đẩy (Push) image lên Docker Hub công khai.\n"
        "5. Chạy thử container và kiểm tra độ phản hồi của API.",
        "tip"
    )

    add_callout(
        doc,
        "Chụp ảnh màn hình khi GitHub Actions báo tích xanh (Passed)",
        "Khi toàn bộ các bước hoàn thành, biểu tượng chuyển sang màu xanh lá cây với dấu tích ✓. Hãy chụp lại màn hình này! Đây là bằng chứng quan trọng nhất khẳng định hệ thống CI/CD đã hoạt động chuẩn xác 100%.",
        "screenshot"
    )

    # SECTION 5
    add_h1("PHẦN 5: KIỂM TRA MÔ HÌNH TRÊN DOCKER HUB VÀ CHẠY THỬ NGHIỆM")
    add_p("Sau khi dây chuyền CD trên GitHub báo thành công (khoảng 3 - 5 phút):")

    add_h2("5.1. Xem Mô Hình Của Bạn Trên Docker Hub")
    add_bullet("Mở trình duyệt web và vào địa chỉ: https://hub.docker.com", "1. ")
    add_bullet("Bấm vào mục 'Repositories'. Bạn sẽ thấy xuất hiện một mục mới mang tên: <tên-của-bạn>/music-cotmix-ai (ví dụ: nimesa0/music-cotmix-ai).", "2. ")
    add_bullet("Bấm vào repository đó, bạn sẽ thấy trạng thái là 'Public' (Công khai) kèm thẻ tag 'latest' và 'sha-...'.", "3. ")

    add_callout(
        doc,
        "Chụp ảnh trang Repository công khai trên Docker Hub",
        "Hãy chụp lại toàn bộ màn hình trang web Docker Hub hiển thị tên Image của bạn cùng thẻ tag 'latest' để làm minh chứng nộp báo cáo.",
        "screenshot"
    )

    add_h2("5.2. Chạy Thử Nghiệm Mô Hình Từ Docker Hub (Dành Cho Mọi Người)")
    add_p(
        "Bây giờ, mô hình AI nhận diện âm nhạc của bạn đã trở thành tài sản công khai trên toàn cầu! "
        "Bất kỳ ai (kể cả khách hàng, giảng viên, hay đồng nghiệp) chỉ cần mở máy tính có Docker và gõ đúng 1 dòng lệnh sau:"
    )
    add_code_block(
        doc,
        "docker run -p 8000:8000 <tên-docker-hub-của-bạn>/music-cotmix-ai:latest\n\n"
        "# Ví dụ cụ thể nếu tên Docker ID của bạn là nimesa0:\n"
        "docker run -p 8000:8000 nimesa0/music-cotmix-ai:latest",
        "Dòng lệnh duy nhất để tải về và chạy mô hình AI ở bất kỳ đâu trên thế giới"
    )

    add_p(
        "Sau đó, mở trình duyệt web và gõ địa chỉ: http://localhost:8000\n"
        "Giao diện web nhận diện thể loại âm nhạc tương tác tuyệt đẹp sẽ lập tức hiển thị. Bạn có thể nghe thử các bài hát mẫu, tải nhạc lên và nhận kết quả phân tích AI trong tích tắc!"
    )

    # SECTION 6
    add_h1("PHẦN 6: PHƯƠNG ÁN DỰ PHÒNG: TỰ BUILD VÀ PUSH BẰNG TAY TRÊN MÁY TÍNH")
    add_p(
        "Nếu sau này bạn muốn tự build và kiểm tra trên máy tính cá nhân trước khi đưa lên mạng, "
        "chúng tôi đã viết sẵn một công cụ tự động 1-click có tên: docker_build_push.bat ngay trong thư mục dự án."
    )
    add_bullet("Tải và cài đặt phần mềm Docker Desktop cho Windows tại: https://www.docker.com/products/docker-desktop/ (nhớ tích chọn 'Use WSL 2').", "Bước 6.1: ")
    add_bullet("Khởi động Docker Desktop và chờ góc dưới bên trái chuyển sang màu xanh lá cây ('Engine running').", "Bước 6.2: ")
    add_bullet("Bấm đúp chuột vào tệp 'docker_build_push.bat' trong thư mục D:\\Music_CoTMix.", "Bước 6.3: ")
    add_bullet("Một cửa sổ màu đen xuất hiện, bạn chỉ cần gõ tên Docker ID của mình và bấm Enter. Phần mềm sẽ tự động đăng nhập, tự động build và tự động đẩy lên Docker Hub!", "Bước 6.4: ")

    # SECTION 7
    add_h1("PHẦN 7: TỔNG HỢP BẢNG CHECKLIST HÀNH ĐỘNG DÀNH CHO BẠN")
    add_p("Dưới đây là bảng tổng hợp các việc bạn cần làm theo đúng thứ tự:")

    chk_tbl = doc.add_table(rows=6, cols=3)
    chk_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(chk_tbl, "CBD5E0")
    
    headers = ["STT", "Hành động cụ thể", "Yêu cầu chụp màn hình (Minh chứng)"]
    for j, h in enumerate(headers):
        c = chk_tbl.cell(0, j)
        set_cell_margins(c, 100, 100, 120, 120)
        set_cell_shading(c, "2B6CB0")
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
    
    chk_tbl.cell(0, 0).width = Inches(0.6)
    chk_tbl.cell(0, 1).width = Inches(3.4)
    chk_tbl.cell(0, 2).width = Inches(2.5)

    checklist_data = [
        ("1", "Đăng ký tài khoản Docker Hub tại hub.docker.com và xác nhận email.", "📸 Chụp màn hình trang chủ Docker Hub sau khi đăng nhập thành công."),
        ("2", "Tạo Access Token trong Account Settings -> Security trên Docker Hub.", "📸 Chụp màn hình hộp thoại Access Token vừa sinh ra."),
        ("3", "Thêm 2 secret (DOCKER_USERNAME và DOCKER_PASSWORD) vào Settings -> Secrets -> Actions trên GitHub.", "📸 Chụp màn hình bảng Repository secrets trên GitHub có đủ 2 secret."),
        ("4", "Chạy lệnh 'git push origin main' hoặc bấm nút 'Run workflow' trong tab Actions trên GitHub.", "📸 Chụp màn hình tiến trình chạy và kết quả tích xanh ✓ của GitHub Actions."),
        ("5", "Kiểm tra kho lưu trữ trên Docker Hub xem đã có image public chưa.", "📸 Chụp màn hình trang repository public trên Docker Hub.")
    ]

    for idx, (stt, act, cap) in enumerate(checklist_data, start=1):
        c0 = chk_tbl.cell(idx, 0)
        c1 = chk_tbl.cell(idx, 1)
        c2 = chk_tbl.cell(idx, 2)
        c0.width = Inches(0.6)
        c1.width = Inches(3.4)
        c2.width = Inches(2.5)
        set_cell_margins(c0, 80, 80, 120, 120)
        set_cell_margins(c1, 80, 80, 120, 120)
        set_cell_margins(c2, 80, 80, 120, 120)
        bg = "FFFFFF" if idx % 2 == 1 else "F7FAFC"
        set_cell_shading(c0, bg)
        set_cell_shading(c1, bg)
        set_cell_shading(c2, bg)
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(stt)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.bold = True
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(act)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(cap)
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(43, 108, 176)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 8
    add_h1("PHẦN 8: XỬ LÝ CÁC VẤN ĐỀ THƯỜNG GẶP (TROUBLESHOOTING)")
    add_bullet(
        " Nguyên nhân do bạn nhập sai tên bí mật (phải gõ đúng chữ in hoa DOCKER_USERNAME và DOCKER_PASSWORD) "
        "hoặc lúc copy Access Token bị thiếu ký tự. Cách khắc phục: Vào GitHub Settings -> Secrets -> Actions, bấm nút hình cây bút để cập nhật lại token.",
        "1. GitHub Actions báo lỗi 'Error response from daemon: login attempt failed':"
    )
    add_bullet(
        " Kiểm tra xem bạn đã lưu tệp .dockerignore mà chúng tôi cung cấp chưa. "
        "Nếu thiếu .dockerignore, hệ thống sẽ cố tải 1.8GB file nhạc mẫu lên gây nghẽn mạng và tốn nhiều thời gian. "
        "Với tệp .dockerignore đã chuẩn bị, thời gian build chỉ mất từ 2 đến 4 phút.",
        "2. Quá trình build trên GitHub Actions chạy quá lâu (> 15 phút):"
    )
    add_bullet(
        " Nguyên nhân do trên máy tính của bạn đang có một chương trình khác chiếm cổng 8000. "
        "Bạn chỉ cần đổi số cổng phía trước sang 8080 hoặc 9000: 'docker run -p 8080:8000 <tên>/music-cotmix-ai:latest' "
        "rồi mở trình duyệt vào http://localhost:8080 là sẽ hoạt động bình thường.",
        "3. Chạy lệnh docker run bị báo lỗi 'port 8000 already in use':"
    )

    # Closing Signature
    add_h2("LỜI KẾT TỪ CHUYÊN GIA PHÁT TRIỂN PHẦN MỀM")
    add_p(
        "Xin chúc mừng bạn! Bằng việc thiết lập quy trình CI/CD và đóng gói Docker Image thành công, dự án trí tuệ nhân tạo "
        "Music_CoTMix của bạn đã chính thức đạt tiêu chuẩn kỹ thuật chuyên nghiệp (Production-Ready MLOps). "
        "Giờ đây, bạn hoàn toàn tự tin mang sản phẩm này đi báo cáo, thuyết trình hoặc triển khai thực tế trên bất kỳ hạ tầng đám mây nào (Render, AWS, GCP, Azure)."
    )

    docx_path = r"D:\Music_CoTMix\Huong_Dan_Trien_Khai_CICD_va_DockerHub_Music_CoTMix.docx"
    pdf_path = r"D:\Music_CoTMix\Huong_Dan_Trien_Khai_CICD_va_DockerHub_Music_CoTMix.pdf"
    doc.save(docx_path)
    print(f"Saved docx to: {docx_path}")
    
    # Export to PDF via Word COM
    print("Exporting docx to PDF via Microsoft Word...")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc_com = word.Documents.Open(docx_path)
        doc_com.ExportAsFixedFormat(pdf_path, 17) # 17 = wdExportFormatPDF
        doc_com.Close(False)
        print(f"Successfully generated PDF at: {pdf_path}")
    finally:
        word.Quit()

if __name__ == "__main__":
    build_pdf_document()
